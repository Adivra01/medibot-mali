from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
import json

base = Path(__file__).parent
metrics = json.loads((base / 'resultats_serieux/metrics_synthea.json').read_text())
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.75); sec.bottom_margin = Inches(.75)
sec.left_margin = Inches(.85); sec.right_margin = Inches(.85)
doc.styles['Normal'].font.name = 'Aptos'; doc.styles['Normal'].font.size = Pt(10.5)
title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('Documentation technique du projet Medibot Mali')
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run('Version Synthea et interface de démonstration').italic = True
doc.add_paragraph('Ce document décrit la version expérimentale séparée du projet Medibot Mali : données, modèle, métriques, interface et limites. Le modèle historique du dossier model_memoire reste séparé.')

def heading(text, level=1): doc.add_heading(text, level=level)
def para(text): doc.add_paragraph(text)
def table(rows):
    t = doc.add_table(rows=1, cols=2); t.style = 'Table Grid'
    t.rows[0].cells[0].text = 'Élément'; t.rows[0].cells[1].text = 'Valeur'
    for a,b in rows:
        c=t.add_row().cells; c[0].text=str(a); c[1].text=str(b)

heading('1. Résumé exécutif')
para('La version Synthea utilise des dossiers médicaux synthétiques pour prédire plusieurs conditions simultanément à partir de caractéristiques patient et d’observations médicales. Il s’agit d’un prototype de recherche et d’aide à l’orientation, pas d’un dispositif de diagnostic clinique.')
heading('2. Nature du modèle')
para('Ce projet n’utilise pas de fine tuning. Aucun modèle de langage pré-entraîné comme BERT ou ClinicalBERT n’a été adapté. Le modèle est un apprentissage supervisé classique : un Random Forest encapsulé dans MultiOutputClassifier.')
para('Le Random Forest construit plusieurs arbres de décision. La sortie est multi-étiquette : un patient peut posséder plusieurs conditions simultanément. Le modèle ne comprend pas directement le langage naturel.')
heading('3. Dataset Synthea')
table([('Patients',111),('Observations',53961),('Conditions',3387),('Variables du modèle',176),('Conditions cibles',10),('Patients de test',28)])
para('Les dix cibles sont les conditions les plus fréquentes avec au moins dix occurrences. Les conditions rares ont été exclues pour éviter une évaluation artificielle.')
heading('4. Préparation des données')
para('Les variables comprennent l’âge, le genre, la race et les observations numériques. Pour chaque patient et chaque observation numérique, la moyenne des valeurs disponibles est calculée. Les textes non numériques ne sont pas transformés artificiellement en symptômes.')
heading('5. Entraînement')
table([('Algorithme','Random Forest multi-étiquette'),('Arbres',300),('Profondeur maximale',8),('Taille minimale des feuilles',2),('Part test','25 %'),('Graine',42)])
para('Le découpage sépare les patients d’entraînement et de test. Les données d’un même patient ne doivent pas apparaître dans les deux ensembles afin d’éviter une fuite d’information.')
heading('6. Résultats')
table([('F1 macro',f"{metrics['f1_macro']*100:.2f} %"),('Hamming loss',f"{metrics['hamming_loss']*100:.2f} %"),('Patients test',metrics['test_patients']),('Accuracy classique','Non retenue comme métrique principale')])
para('L’accuracy classique n’est pas la métrique principale pour une sortie multi-étiquette : elle exige que toutes les conditions soient correctes en même temps. Le F1 macro mesure l’équilibre entre les conditions ; le Hamming loss mesure la proportion de décisions condition-étiquette incorrectes.')
para('Les résultats sont indicatifs : l’échantillon ne contient que 111 patients et Synthea produit des données synthétiques. Ils ne prouvent pas une performance clinique réelle.')
heading('7. Interface')
para('L’application Streamlit utilise un fond clair. Les symptômes fréquents sont affichés sous forme de boutons visibles. Un champ de recherche permet l’autocomplétion dans le vocabulaire du dataset. L’utilisateur reçoit trois résultats potentiels, leur pourcentage, un niveau d’urgence indicatif et une recommandation générale. Les consultations sont enregistrées localement dans SQLite.')
para('Un symptôme absent du vocabulaire est signalé mais n’est pas injecté dans le modèle. Le réentraînement automatique serait incorrect sans exemples médicaux étiquetés ; les nouveaux symptômes doivent être validés et documentés avant tout réentraînement.')
heading('8. Exécution')
para('Depuis model_memoire_synthea :')
para('source .venv/bin/activate\npython 04_modele_synthea.py\nstreamlit run app.py')
para('Les sorties sont metrics_synthea.json, classification_report_synthea.txt et random_forest_synthea_multilabel.pkl.')
heading('9. Limites et sécurité')
para('Le système ne doit pas être présenté comme un dispositif médical certifié. Les probabilités du Random Forest ne sont pas des probabilités cliniques calibrées. Une utilisation hospitalière nécessiterait des données réelles validées, une revue médicale, une étude prospective, une calibration, des tests de biais, une gestion de l’incertitude et une conformité réglementaire.')
para('Le niveau d’urgence de l’interface est une règle de démonstration et non une décision médicale. Toute situation grave doit conduire à contacter immédiatement les services de santé appropriés.')
heading('10. Conclusion')
para('La version Synthea est un modèle Random Forest multi-étiquette, et non un fine tuning. Elle démontre une chaîne reproductible sur des dossiers synthétiques. Le modèle symptomatique historique reste séparé pour assurer la continuité du mémoire et permettre une comparaison future.')
out = base / 'Documentation_technique_Medibot_Mali.docx'; doc.save(out); print(out)
